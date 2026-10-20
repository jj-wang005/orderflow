"""Bounded local input readers. Never execute formulas, fetch URLs or call a model."""
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import PureWindowsPath
import base64
import json
import warnings
from app.api.errors import APIError
from app.schemas.ai_extraction import ExtractionResult, ExtractionRow
from app.services.excel_parser import _read_source, _check_actual_bounds, ParseLimits
from app.services.customer_management import ALIASES
from app.services.data_normalizer import normalize_code, normalize_quantity


@dataclass
class InputDocument:
    sources: list[dict] = field(default_factory=list)
    images: dict[str, str] = field(default_factory=dict)
    program_result: ExtractionResult | None = None


def image_url(content):
    from PIL import Image, ImageOps
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('error', Image.DecompressionBombWarning)
            with Image.open(BytesIO(content)) as image:
                if image.format not in {'JPEG','PNG'} or image.width*image.height > 12_000_000:
                    raise ValueError()
                image = ImageOps.exif_transpose(image).convert('RGB')
                image.thumbnail((2400,2400))
                buffer=BytesIO();image.save(buffer,format='JPEG',quality=85)
                return 'data:image/jpeg;base64,'+base64.b64encode(buffer.getvalue()).decode()
    except Exception:
        raise APIError(422,'INVALID_IMAGE','图片损坏、格式不支持或像素超过 1200 万') from None


def _excel(content,name):
    from openpyxl import load_workbook
    limits=ParseLimits(max_rows=500,max_columns=50,max_uncompressed_bytes=20*1024*1024)
    _read_source(content,name,limits)
    book=load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
    doc=InputDocument();results=[];stable=True
    try:
        if len(book.sheetnames)>5:raise ValueError('最多 5 张工作表')
        for si,sheet in enumerate(book):
            _check_actual_bounds(content,sheet.title,1,limits);sheet.reset_dimensions()
            rows=list(sheet.iter_rows(max_row=501,max_col=51,values_only=True))
            nonempty=[(i+1,list(values)) for i,values in enumerate(rows) if any(v is not None for v in values)]
            if not nonempty:continue
            headers=[]
            for pos,values in nonempty[:30]:
                mapping={}
                for ci,value in enumerate(values):
                    if not isinstance(value,str):continue
                    matches=[k for k,aliases in ALIASES.items() if value.strip().casefold() in {a.casefold() for a in aliases|{k}}]
                    if len(matches)==1:mapping[ci]=matches[0]
                if 'quantity' in mapping.values() and len(set(mapping.values()))>=2 and len(set(mapping.values()))==len(mapping):headers.append((pos,mapping))
            if len(headers)!=1:stable=False
            for pos,values in nonempty:
                sid=f'xlsx:{si+1}:{pos}'
                text=' | '.join(f'{i+1}={v}' for i,v in enumerate(values) if v is not None)
                doc.sources.append({'id':sid,'location':f'{sheet.title} 第 {pos} 行','text':text,'visual':False})
                if len(headers)==1 and pos>headers[0][0]:
                    data={f:None for f in ('customer_order_no','customer_material_code','customer_model','customer_batch','quantity')}
                    try:
                        for index,target in headers[0][1].items():
                            value=values[index]
                            if isinstance(value,str) and value.startswith('='):raise ValueError()
                            data[target]=str(normalize_quantity(value)) if target=='quantity' else normalize_code(value)
                        results.append(ExtractionRow(**data,source_id=sid,source_excerpt=text[:2000],extraction_warnings=[],uncertain_fields=[]))
                    except ValueError:stable=False
        if stable and results:doc.program_result=ExtractionResult(rows=results,warnings=['程序识别表头；仍请核对并确认'])
        return doc
    finally:book.close()


def prepare_input(kind,content,name,max_pages=5):
    if not content or len(content)>10*1024*1024:raise APIError(413,'INPUT_TOO_LARGE','输入为空或超过 10 MiB')
    try:
        if kind=='TEXT':
            value=content.decode('utf-8')
            if not value.strip() or len(value)>30000:raise ValueError()
            doc=InputDocument(sources=[{'id':'text:1','location':'粘贴文字','text':value,'visual':False}])
        elif kind=='XLSX':doc=_excel(content,name)
        elif kind=='IMAGE':doc=InputDocument(sources=[{'id':'image:1','location':'原始图片','text':'','visual':True}],images={'image:1':image_url(content)})
        elif kind=='PDF':
            import pymupdf
            doc=InputDocument()
            with pymupdf.open(stream=content,filetype='pdf') as pdf:
                if pdf.needs_pass or not 1<=pdf.page_count<=max_pages:raise APIError(422,'PDF_LIMIT','PDF 加密、无页面或超过允许页数')
                for i,page in enumerate(pdf):
                    sid=f'pdf:{i+1}';text=page.get_text('text',sort=True)
                    visual=len(text.strip())<20
                    # Include any image-bearing page as well: mixed scans must not silently disappear.
                    visual=visual or bool(page.get_images())
                    tables=[]
                    if text.strip():
                        for table in page.find_tables().tables:
                            values=table.extract()
                            if len(tables)>=20 or len(values)>100 or any(len(r)>50 for r in values):
                                raise APIError(422,'INPUT_LIMIT','PDF 表格过多或过大，请拆分后解析')
                            tables.append(values)
                        if tables:text+='\n提取表格：'+json.dumps(tables,ensure_ascii=False)
                    if visual:
                        if page.rect.width*page.rect.height>8_000_000:raise APIError(422,'PDF_LIMIT','PDF 页面尺寸超过限制')
                        scale=min(1.6,2400/max(page.rect.width,page.rect.height))
                        doc.images[sid]=image_url(page.get_pixmap(matrix=pymupdf.Matrix(scale,scale),alpha=False).tobytes('png'))
                    doc.sources.append({'id':sid,'location':f'PDF 第 {i+1} 页','text':text,'visual':visual,'tables':tables})
        else:raise ValueError()
        if not doc.sources or sum(len(s['text']) for s in doc.sources)>30000:raise APIError(422,'INPUT_LIMIT','输入无内容或提取文字超过 30000 字，请拆分文件')
        if doc.program_result and len(doc.program_result.rows)>100:raise APIError(422,'ROW_LIMIT','明细超过 100 行，请使用标准 Excel 或拆分')
        return doc
    except APIError:raise
    except Exception:raise APIError(422,'INPUT_INVALID','文件无法安全解析，请检查格式、工作表范围或改用粘贴文字') from None


def source_name(name):
    result=PureWindowsPath(name or '').name
    if not result or len(result)>255 or any(ord(c)<32 for c in result):raise APIError(422,'INVALID_FILE','文件名不合法')
    return result
