"""Read-only internal exports; only explicitly selected business fields leave DB."""
from io import BytesIO
from decimal import Decimal
import fitz
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from sqlalchemy import select
from app.models import SalesOrder,CustomerConfig,OrderChange,ImportTaskRow
from app.repositories.api_queries import order_conditions,get_order
from app.api.errors import APIError

MAX_EXPORT_ROWS=10000

def excel_orders(session,**filters):
    # One bounded query avoids count/page drift, and uses the exact list predicate.
    pairs=session.execute(select(SalesOrder,CustomerConfig.customer_name).join(CustomerConfig,SalesOrder.customer_id==CustomerConfig.customer_id)
        .where(*order_conditions(**filters)).order_by(SalesOrder.created_at.desc(),SalesOrder.id.desc()).limit(MAX_EXPORT_ROWS+1)).all()
    if len(pairs)>MAX_EXPORT_ROWS:raise APIError(422,'EXPORT_TOO_LARGE','导出超过10000条，请缩小筛选日期或客户范围')
    book=Workbook();sheet=book.active;sheet.title='销售订单'
    sheet.append(['销售订单编号','客户编号','客户名称','客户采购单号','客户料号','客户机型','客户批次','我司产品编码','我司产品名称','数量','销售单价','订单金额','状态','创建时间','建单方式','来源任务','来源工作表','来源行','历史参考订单ID'])
    from app.schemas.order_status import STATUS_LABELS
    for order,name in pairs:
        detail=get_order_description(order)
        values=[order.order_no,order.customer_id,name,order.customer_order_no,order.customer_material_code,order.customer_model,order.customer_batch,order.product_code,order.product_name,
                order.quantity,order.unit_price,order.amount,STATUS_LABELS.get(order.status,order.status),order.created_at,detail,order.import_task_id,order.source_sheet_name,order.source_row_number,order.source_order_id]
        # Excel numbers retain only 15 significant digits; preserve larger DB decimals as text.
        sheet.append([ILLEGAL_CHARACTERS_RE.sub('',v) if isinstance(v,str) else str(v) if isinstance(v,Decimal) and len(v.normalize().as_tuple().digits)>15 else v for v in values])
        for cell in sheet[sheet.max_row]:
            # Force customer-provided strings to OOXML string cells, including '=' / tabs.
            if isinstance(cell.value,str):cell.data_type='s'
        for col,fmt in [(10,'0.###'),(11,'0.0000'),(12,'#,##0.00'),(14,'yyyy-mm-dd hh:mm:ss')]:sheet.cell(sheet.max_row,col).number_format=fmt
    from openpyxl.utils import get_column_letter
    widths=[40,12,24,28,26,24,24,26,32,14,16,18,16,23,22,14,22,12,20]
    for index,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(index)].width=width
    for cell in sheet[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='244568');cell.alignment=Alignment(vertical='center')
    sheet.row_dimensions[1].height=28;sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    buffer=BytesIO();book.save(buffer);return buffer.getvalue()

def get_order_description(order):
    if order.source_type=='SEED':return '模拟历史参考'
    if order.is_simulated_manual:return '人工补全建单'
    return '历史匹配建单' if order.source_order_id else '来源待核实'

def pdf_order(session,order_id):
    detail=get_order(session,order_id)
    customer=session.scalar(select(CustomerConfig.customer_name).where(CustomerConfig.customer_id==detail.customer_id))
    logs=list(session.scalars(select(OrderChange).where(OrderChange.order_id==order_id).order_by(OrderChange.id).limit(1001)))
    if len(logs)>1000:raise APIError(422,'EXPORT_TOO_LARGE','订单变更记录过多，请联系管理员分段核对')
    row=session.scalar(select(ImportTaskRow).where(ImportTaskRow.created_order_id==order_id).order_by(ImportTaskRow.id))
    doc=fitz.open();font=fitz.Font('china-s');page=None;y=0
    def new_page():
        nonlocal page,y
        page=doc.new_page(width=595,height=842)
        page.insert_font(fontname='cvte-cn',fontbuffer=font.buffer)
        page.insert_text((42,43),'OrderFlow / 内部销售订单详情',fontname='cvte-cn',fontsize=17,color=(.12,.24,.36))
        page.insert_text((42,65),'内部查询核对报告 - 非正式 OMS 单据、合同或发票',fontname='cvte-cn',fontsize=10,color=(.4,.4,.4))
        page.draw_line((42,78),(553,78),color=(.75,.8,.85));y=105
    def line(value,size=11):
        nonlocal y
        # Wrap by measured glyph width; controls cannot affect document structure.
        value=''.join(c if ord(c)>=32 else ' ' for c in str(value));segment=''
        for c in value:
            if font.text_length(segment+c,fontsize=size)>505:
                if y>785:new_page()
                page.insert_text((42,y),segment,fontname='cvte-cn',fontsize=size);y+=19;segment=''
            segment+=c
        if y>785:new_page()
        page.insert_text((42,y),segment,fontname='cvte-cn',fontsize=size);y+=21
    def field(label,value):line(label+'：'+(str(value) if value is not None else '未提供'))
    new_page()
    field('销售订单编号',detail.order_no);field('状态',detail.status_description);field('客户',f'{customer or detail.customer_id} ({detail.customer_id})')
    for label,key in [('客户采购单号','customer_order_no'),('客户料号','customer_material_code'),('客户机型','customer_model'),('客户批次','customer_batch'),('我司产品编码','product_code'),('我司产品名称','product_name'),('采购数量','quantity'),('销售单价','unit_price'),('后端保存总金额','amount')]:field(label,getattr(detail,key))
    field('建单方式',detail.creation_description);field('技术来源',detail.source_type)
    field('导入任务',detail.import_task_id);field('采购行ID',row.id if row else None);field('来源工作表',detail.source_sheet_name);field('来源行',detail.source_row_number)
    if detail.source_order_id:
        reference=session.get(SalesOrder,detail.source_order_id)
        if reference:field('历史参考订单',f'#{reference.id} / {reference.order_no}')
    field('创建时间',detail.created_at);field('产品价格说明',detail.price_description)
    line('状态及变更审计',13)
    if not logs:line('无系统变更记录')
    for log in logs:
        field('操作时间',log.created_at);field('操作来源',log.actor_label)
        before=(log.before_data or {}).get('status','新建');after=log.after_data.get('status','')
        field('状态记录',before+' → '+after);field('作废原因' if log.target.startswith('void:') else '说明',log.note)
    for number,p in enumerate(doc,1):p.insert_text((42,817),f'OrderFlow 内部核对 / 第 {number} 页，共 {len(doc)} 页',fontname='cvte-cn',fontsize=9,color=(.4,.4,.4))
    doc.subset_fonts()
    result=doc.tobytes(garbage=4,deflate=True);doc.close();return result
