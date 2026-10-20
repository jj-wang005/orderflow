"""One current configuration, immutable version receipts, no sales-order writes."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select, func
from sqlalchemy.exc import IntegrityError
from app.models import CustomerConfig, MatchRule, ConfigurationVersion, ImportTask
from app.schemas.customer_configuration import CustomerConfiguration
from app.schemas.customer_management import ConfigurationSave
from app.schemas.order_status import HISTORICAL_STATUSES, HISTORICAL_SOURCES
from app.services.excel_parser import _read_source, _check_actual_bounds, _json_value, ParseLimits
from app.api.errors import APIError
from app.services.data_normalizer import normalize_code
from app.services.source_columns import configuration_columns

FIELD_LABELS = {"customer_order_no": "客户采购单号", "customer_material_code": "客户料号", "customer_batch": "客户批次", "customer_model": "客户机型", "quantity": "采购数量"}
ALIASES = {
    # Exact headers only (case-insensitive). Keep ambiguous names such as
    # 型号 / 客户/型号名称, internal product codes and stock quantities out.
    # Used by suggestions only; saved customer mappings remain authoritative.
    "customer_order_no": {"采购订单号", "客户采购单号", "客户订单号", "采购单号", "单号", "PO", "PO号",
                          "客户单号", "客户采购订单号", "采购订单编号", "PO编号", "PO No", "PO No.", "Purchase Order No", "Purchase Order No."},
    "customer_material_code": {"物料/品项编码", "物料编码", "客户料号", "料号", "客户物料编码",
                               "物料编号", "物料代码", "品项编码", "客户物料号", "客户物料编号", "客户物料代码",
                               "Customer Material Code", "Customer Part No", "Customer Part No."},
    "customer_batch": {"客批号", "客户批次", "客户批号", "批次", "批号", "客户批次号", "客户批次编号", "Customer Batch"},
    "customer_model": {"机型", "客户机型", "客户型号", "客户机型编码", "Customer Model"},
    "quantity": {"数量", "订单数量", "采购数量", "本次数量", "qty", "订购数量", "订货数量", "采购订单数量",
                 "Quantity", "Order Qty", "Order Quantity", "Ordered Quantity"},
}
TEMPLATES = [
    {"id": "A", "name": "采购单号 + 料号一致", "match_fields": ["customer_order_no", "customer_material_code"], "historical_quantity": None},
    {"id": "B", "name": "采购单号 + 客户机型", "match_fields": ["customer_order_no", "customer_model"], "historical_quantity": None},
    {"id": "C", "name": "客户批次 + 料号，历史数量 > 500", "match_fields": ["customer_batch", "customer_material_code"], "historical_quantity": {"operator": "gt", "value": "500"}},
]

def customer_record(session, customer_id, lock=False):
    query = select(CustomerConfig).where(CustomerConfig.customer_id == customer_id).execution_options(populate_existing=True)
    customer = session.scalar(query.with_for_update() if lock else query)
    if customer is None:
        raise APIError(404, "CUSTOMER_NOT_FOUND", "客户不存在")
    return customer

def configuration_data(session, customer):
    rule = session.scalar(select(MatchRule).where(MatchRule.customer_id == customer.customer_id, MatchRule.enabled.is_(True)))
    if rule is None:
        return None
    return {"source_columns": configuration_columns(session, customer)[0],
            "customer": {k: getattr(customer, k) for k in CustomerConfiguration.model_fields["customer"].annotation.model_fields},
            "rule": {k: getattr(rule, k) for k in CustomerConfiguration.model_fields["rule"].annotation.model_fields}}

def detail(session, customer_id):
    customer = customer_record(session, customer_id)
    config = configuration_data(session, customer)
    versions = session.scalars(select(ConfigurationVersion).where(ConfigurationVersion.customer_id == customer_id).order_by(ConfigurationVersion.config_version.desc())).all()
    return {"customer_id": customer_id, "customer_name": customer.customer_name, "enabled": customer.enabled,
            "config_version": customer.config_version, "configured": config is not None, "configuration": config,
            "header_source": configuration_columns(session, customer)[1] if config else 'upload_required',
            "versions": [{"config_version": v.config_version, "rule_version": v.rule_version, "created_at": v.created_at,
                          "snapshot": v.snapshot, "task_count": session.scalar(select(func.count()).select_from(ImportTask).where(ImportTask.customer_id == customer_id, ImportTask.config_version == v.config_version))} for v in versions]}

def list_customers(session):
    rules = set(session.scalars(select(MatchRule.customer_id).where(MatchRule.enabled.is_(True))))
    return [{"customer_id": c.customer_id, "customer_name": c.customer_name, "enabled": c.enabled,
             "config_version": c.config_version, "configured": c.customer_id in rules}
            for c in session.scalars(select(CustomerConfig).order_by(CustomerConfig.customer_id))]

def create_customer(session, request):
    if session.scalar(select(CustomerConfig.id).where(CustomerConfig.customer_id == request.customer_id)) is not None:
        raise APIError(409, "CUSTOMER_EXISTS", "客户标识已存在，请编辑原客户")
    try:
        with session.begin_nested():
            session.add(CustomerConfig(customer_id=request.customer_id, customer_name=request.customer_name,
                        sheet_name="Sheet1", header_row=1, field_mapping={}, enabled=False, config_version=0))
            session.flush()
    except IntegrityError:
        raise APIError(409, "CUSTOMER_EXISTS", "客户标识已存在") from None
    return detail(session, request.customer_id)

def build_configuration(customer_id, body, version):
    fields = list(body.match_fields)
    validation = {"extra_required_fields": [], "consistency_fields": []}
    # An explicit rule save replaces the selection for every customer equally.
    # Legacy safety fields remain in existing rules, version receipts and tasks.
    try:
        return CustomerConfiguration.model_validate({
            "customer": {"customer_id": customer_id, "customer_name": body.customer_name, "sheet_name": body.sheet_name,
                         "header_row": body.header_row, "field_mapping": body.field_mapping, "enabled": body.enabled, "config_version": version},
            "rule": {"match_fields": fields, "filter_conditions": {"customer_scope": "same_customer", "eligible_statuses": HISTORICAL_STATUSES,
                     "allowed_source_types": HISTORICAL_SOURCES, "historical_quantity": body.historical_quantity.model_dump(mode="json") if body.historical_quantity else None},
                     "sort_config": [{"field": "created_at", "direction": "DESC"}, {"field": "id", "direction": "DESC"}],
                     "validation_config": validation, "enabled": True, "rule_version": version},
            "source_columns": [c.model_dump() for c in (body.source_columns or [])]})
    except ValidationError:
        raise APIError(422, "INVALID_CONFIGURATION", "配置不合法：数量及匹配字段必须有唯一映射，匹配字段不可重复；请核对表头与工作表") from None

def archive(session, config):
    cid, version = config["customer"]["customer_id"], config["customer"]["config_version"]
    if session.scalar(select(ConfigurationVersion.id).where(ConfigurationVersion.customer_id == cid, ConfigurationVersion.config_version == version)) is None:
        session.add(ConfigurationVersion(customer_id=cid, config_version=version, rule_version=config["rule"]["rule_version"], snapshot=config))

def save_configuration(session, customer_id, body, *, preserve_rule=False):
    customer = customer_record(session, customer_id, lock=True)
    if body.expected_version != customer.config_version:
        raise APIError(409, "CONFIG_VERSION_CONFLICT", "配置已被其他操作修改，请重新打开后核对")
    rules = session.scalars(select(MatchRule).where(MatchRule.customer_id == customer_id).with_for_update()).all()
    version = max([customer.config_version, *[r.rule_version for r in rules]]) + 1
    old_columns = configuration_columns(session, customer)[0]
    if body.source_columns is None:
        if old_columns and (body.sheet_name != customer.sheet_name or body.header_row != customer.header_row):
            raise APIError(422, 'HEADERS_REQUIRED', '工作表或表头行已改变，请重新上传样例读取全部表头')
        body = ConfigurationSave.model_validate({**body.model_dump(), 'source_columns': old_columns})
    elif not body.source_columns and old_columns:
        raise APIError(422, 'HEADERS_REQUIRED', '不能清空已保存的原始表头，请重新上传样例')
    config = build_configuration(customer_id, body, version)
    previous = configuration_data(session, customer)
    if preserve_rule and previous:
        # Basic customer edits must not discard legacy extra required fields,
        # consistency checks or historical filters. Only explicit rule saves do.
        config = CustomerConfiguration.model_validate({
            **config.model_dump(mode="json"),
            "rule": {**previous["rule"], "rule_version": version}})
    if previous:
        archive(session, previous)
    for key, value in config.customer.model_dump(mode="json").items():
        setattr(customer, key, value)
    for rule in rules:
        rule.enabled = False
    session.add(MatchRule(customer_id=customer_id, **config.rule.model_dump(mode="json")))
    archive(session, config.model_dump(mode="json"))
    session.flush()
    return detail(session, customer_id)

def edit_customer(session, customer_id, body):
    customer = customer_record(session, customer_id, lock=True)
    if body.expected_version != customer.config_version:
        raise APIError(409, "CONFIG_VERSION_CONFLICT", "配置已变化，请重新加载")
    config = configuration_data(session, customer)
    if config:
        rule = config["rule"]
        return save_configuration(session, customer_id, ConfigurationSave(
            expected_version=body.expected_version, customer_name=body.customer_name, sheet_name=customer.sheet_name,
            header_row=customer.header_row, field_mapping=customer.field_mapping, enabled=customer.enabled,
            match_fields=list(dict.fromkeys(rule["match_fields"] + rule["validation_config"].get("consistency_fields", []))),
            historical_quantity=rule["filter_conditions"].get("historical_quantity")), preserve_rule=True)
    customer.customer_name = body.customer_name
    customer.config_version += 1
    session.flush()
    return detail(session, customer_id)

def suggest_mapping(content, filename, sheet_name=None, header_row=1):
    limits = ParseLimits()
    book = None
    try:
        _, content = _read_source(content, filename, limits)
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        sheets = [s for s in book.sheetnames if s not in {"字段映射", "字段映射说明"}]
        selected = sheet_name or (sheets[0] if sheets else "")
        _check_actual_bounds(content, selected, header_row, limits)
        ws = book[selected]
        ws.reset_dimensions()
        rows = ws.iter_rows(min_row=header_row, max_row=header_row + 5, max_col=limits.max_columns + 1)
        headers = next(rows, ())
        if any(c.value is not None for c in headers[limits.max_columns:]):
            raise ValueError()
        headers = list(headers[:limits.max_columns])
        samples = list(rows)
        while headers and headers[-1].value is None and not any(len(r)>=len(headers) and r[len(headers)-1].value is not None for r in samples):
            headers.pop()
        names = [c.value for c in headers]
        if not names or any(n is not None and not isinstance(n, str) for n in names) or any(c.data_type == "f" for c in headers):
            raise APIError(422, "INVALID_HEADERS", "表头需为文本，不支持公式或数字表头；请核对表头行")
        columns = []
        for index, name in enumerate(names):
            cells = [r[index] for r in samples if index < len(r) and r[index].value is not None]
            values = [_json_value(c.value) for c in cells]
            mappable = isinstance(name,str) and bool(name.strip()) and name == name.strip() and names.count(name) == 1
            candidates = [f for f, aliases in ALIASES.items() if mappable and name.casefold() in {a.casefold() for a in aliases | {f}}]
            target = candidates[0] if len(candidates) == 1 else None
            if any(c.data_type in {"f", "e", "d"} or isinstance(c.value, bool) for c in cells):
                target = None
            if target and target != "quantity":
                try:
                    for cell in cells:
                        normalize_code(cell.value, cell.number_format)
                except ValueError:
                    target = None
            if target == "quantity":
                try:
                    numbers = [Decimal(str(c.value)) for c in cells]
                    if not numbers or any(not n.is_finite() or n <= 0 or n != n.to_integral_value() or n >= Decimal("1e15") for n in numbers):
                        target = None
                except (InvalidOperation, ValueError):
                    target = None
            columns.append({"index": index+1, "column": name, "mappable": mappable, "samples": values, "recommended_field": target,
                            "reason": ('空白、重复或带首尾空格的表头：保留位置，暂不映射；需要映射时请先整理原表头' if not mappable else
                                       '表头别名及样例类型一致' if target else '未映射，请按实际业务选择')})
        for column in columns:
            target = column["recommended_field"]
            if target and sum(c["recommended_field"] == target for c in columns) > 1:
                # Mark all colliding columns before mutating recommendations.
                column["ambiguous"] = True
        for column in columns:
            if column.pop("ambiguous", False):
                column.update(recommended_field=None, reason="多个列可能对应同一字段，请人工选择")
        mapped = {c["recommended_field"] for c in columns if c["recommended_field"]}
        return {"sheet_names": sheets, "sheet_name": selected, "header_row": header_row, "columns": columns,
                "templates": [{**t, "recommended": set(t["match_fields"]) <= mapped} for t in TEMPLATES]}
    except APIError:
        raise
    except Exception as exc:
        # File input never exposes XML/parser internals or server paths.
        raise APIError(422, "INVALID_SAMPLE", "样例无法读取或超出限制，请使用合法 XLSX 并核对工作表、表头行") from None
    finally:
        if book is not None:
            book.close()
