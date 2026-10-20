"""Read-only XLSX parsing driven by validated module 02 configuration."""
import logging
import math
import posixpath
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError, fromstring, iterparse
from zlib import error as ZlibError
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, ConfigDict, Field, ValidationError
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session
from app.repositories.customer_configuration import load_configuration, ConfigurationError
from app.schemas.customer_configuration import CustomerConfiguration
from app.schemas.purchase_row import FileParseResult, ParseIssue, RawPurchaseRow
from app.services.data_normalizer import normalize_row, is_blank

logger = logging.getLogger(__name__)

class ParseLimits(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    max_size_bytes: int = Field(default=10 * 1024 * 1024, gt=0)
    max_uncompressed_bytes: int = Field(default=100 * 1024 * 1024, gt=0)
    max_rows: int = Field(default=10000, gt=0)
    max_columns: int = Field(default=256, gt=0, le=16384)

class FileProblem(ValueError):
    def __init__(self, code: str, message: str, fields: list[str] | None = None):
        self.issue = ParseIssue(code=code, message=message, fields=fields or [])
        super().__init__(message)

def _json_value(value):
    # JSON-safe originals for later import_task_row.raw_data; format/type kept separately.
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    # openpyxl array formula values expose text without executing the formula.
    return str(getattr(value, "text", value))

def _read_source(source: Path | str | bytes, file_name: str | None, limits: ParseLimits):
    if isinstance(source, bytes):
        name = Path(file_name or "").name
        content = source
    elif isinstance(source, (str, Path)):
        path = Path(source)
        name = path.name
        with path.open("rb") as handle:
            content = handle.read(limits.max_size_bytes + 1)
    else:
        raise FileProblem("INVALID_FILE", "输入必须为本地路径或带文件名的 bytes")
    if Path(name).suffix.lower() != ".xlsx":
        raise FileProblem("INVALID_FILE", "只支持 .xlsx 文件；bytes 输入必须提供 file_name")
    if not content or len(content) > limits.max_size_bytes:
        raise FileProblem("INVALID_FILE", "文件为空或超过大小限制")
    with ZipFile(BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(e.file_size for e in entries) > limits.max_uncompressed_bytes:
            raise FileProblem("INVALID_FILE", "XLSX 解压内容超过限制")
        if any(e.flag_bits & 1 for e in entries) or "xl/vbaProject.bin" in archive.namelist():
            raise FileProblem("INVALID_FILE", "不支持加密或含宏工作簿")
        if len({e.filename for e in entries}) != len(entries):
            raise FileProblem("INVALID_FILE", "XLSX 包内存在重复文件名")
    return name, content

def _check_actual_bounds(content, sheet_name, header_row, limits):
    # Validate actual XML positions as well as producer dimensions; do not silently
    # truncate cells outside max_col when a malformed file advertises A1:A1.
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    relationship_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with ZipFile(BytesIO(content)) as archive:
        sheets = fromstring(archive.read("xl/workbook.xml")).findall("s:sheets/s:sheet", ns)
        sheet = next((item for item in sheets if item.get("name") == sheet_name), None)
        if sheet is None:
            raise FileProblem("INVALID_FILE", "配置指定的工作表不存在", [sheet_name])
        relation_id = sheet.get("{" + relationship_ns + "}id")
        relations = fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        relation = next((item for item in relations if item.get("Id") == relation_id), None)
        if relation is None or relation.get("TargetMode") == "External":
            raise FileProblem("INVALID_FILE", "工作表引用不合法")
        target = relation.get("Target", "")
        target = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
        if not target.startswith("xl/worksheets/"):
            raise FileProblem("INVALID_FILE", "指定对象不是数据工作表")
        with archive.open(target) as xml:
            for _, element in iterparse(xml, events=("end",)):
                tag = element.tag.rsplit("}", 1)[-1]
                if tag == "row":
                    if int(element.get("r", "0")) > header_row + limits.max_rows:
                        raise FileProblem("INVALID_FILE", "工作表实际行数超过限制")
                elif tag == "c":
                    coordinate = element.get("r", "")
                    letters = coordinate.rstrip("0123456789")
                    if letters and column_index_from_string(letters) > limits.max_columns:
                        raise FileProblem("INVALID_FILE", "工作表实际列数超过限制")
                element.clear()

def _parse_workbooks(config, content, result, limits, trust_formula_cache):
    original = cached = None
    try:
        _check_actual_bounds(content, config.customer.sheet_name, config.customer.header_row, limits)
        original = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        sheet_name = config.customer.sheet_name
        if sheet_name not in original.sheetnames:
            raise FileProblem("INVALID_FILE", "配置指定的工作表不存在", [sheet_name])
        sheet, cached_sheet = original[sheet_name], cached[sheet_name]
        start = config.customer.header_row
        if start > limits.max_rows:
            raise FileProblem("CONFIG_ERROR", "表头位置超过解析行数限制")
        for ws in (sheet, cached_sheet):
            if ws.max_column and ws.max_column > limits.max_columns:
                raise FileProblem("INVALID_FILE", "工作表列数超过限制")
            if ws.max_row and ws.max_row > start + limits.max_rows:
                raise FileProblem("INVALID_FILE", "工作表行数超过限制")
            # Do not trust too-small producer dimensions: retain actual data rows.
            ws.reset_dimensions()
        rows = sheet.iter_rows(max_col=limits.max_columns + 1)
        cache_rows = cached_sheet.iter_rows(max_col=limits.max_columns + 1)
        headers = None
        mapping = config.customer.field_mapping
        for index, (cells, cache_cells) in enumerate(zip(rows, cache_rows), start=1):
            if index > start + limits.max_rows:
                raise FileProblem("INVALID_FILE", "工作表行数超过限制")
            if not is_blank(cells[-1].value):
                raise FileProblem("INVALID_FILE", "工作表列数超过限制")
            cells, cache_cells = cells[:-1], cache_cells[:-1]
            if index < start:
                continue
            if index == start:
                headers = []
                for cell in cells:
                    value = cell.value
                    if is_blank(value):
                        headers.append(None)
                    elif getattr(cell, "data_type", None) == "f" or not isinstance(value, str):
                        raise FileProblem("INVALID_FILE", "表头必须为明确的文本列名")
                    else:
                        headers.append(value)
                names = [h for h in headers if h is not None]
                duplicates = sorted({h for h in names if names.count(h) > 1})
                if set(duplicates) & set(mapping):
                    raise FileProblem("INVALID_FILE", "表头重复，不能覆盖原始数据", duplicates)
                required_columns = [column for column, field in mapping.items() if field in config.rule.required_fields]
                missing = [column for column in required_columns if column not in names]
                if missing:
                    raise FileProblem("MISSING_COLUMN", "缺少必要表头：" + ", ".join(missing), missing)
                generated = [f"__column_{get_column_letter(i)}__" for i,h in enumerate(headers,1) if h is None or h in duplicates]
                if set(generated) & set(names):
                    raise FileProblem("INVALID_FILE", "表头与无名列追溯标识冲突")
                continue
            if all(is_blank(cell.value) for cell in cells):
                continue
            raw_data, raw_cells, values, cell_errors = {}, {}, {}, {}
            for col, (header, cell, cache_cell) in enumerate(zip(headers, cells, cache_cells),start=1):
                if header is None and is_blank(cell.value):
                    continue
                key = header if header is not None and header not in duplicates else f"__column_{get_column_letter(col)}__"
                value = cell.value
                data_type = getattr(cell, "data_type", "n")
                raw_data[key] = _json_value(value)
                raw_cells[key] = {"coordinate": f"{get_column_letter(col)}{index}", "data_type": data_type,
                                  "number_format": getattr(cell, "number_format", "General"),
                                  "original_type": type(value).__name__, "original_header": header, "column_index": col}
                if header is None:
                    cell_errors[key] = "无表头列包含数据，已按列坐标保留，请核对模板"
                elif data_type == "f":
                    cache_value = cache_cell.value
                    raw_cells[key]["cached_value"] = _json_value(cache_value)
                    # A cache alone is not proof of freshness. Caller must explicitly trust it.
                    values[key] = cache_value if trust_formula_cache else None
                    if key in mapping and (not trust_formula_cache or is_blank(cache_value) or getattr(cache_cell,"data_type",None)=="e"):
                        cell_errors[key] = "公式字段缺少可信缓存结果；服务端不执行公式"
                elif data_type == "e" and key in mapping:
                    cell_errors[key] = "Excel 单元格包含错误值"
                else:
                    values[key] = value
            raw = RawPurchaseRow(sheet_name=sheet_name, row_number=index, raw_data=raw_data,
                                 raw_cells=raw_cells, values=values, cell_errors=cell_errors)
            result.rows.append(normalize_row(config, raw))
        if headers is None or not result.rows:
            raise FileProblem("INVALID_FILE", "工作表缺少表头或没有采购数据行")
    finally:
        if original is not None:
            original.close()
        if cached is not None:
            cached.close()

def parse_excel(config: CustomerConfiguration, source: Path | str | bytes, *,
                file_name: str | None = None, limits: ParseLimits | None = None,
                trust_formula_cache: bool = False) -> FileParseResult:
    """Pure parsing with an existing configuration. No database writes/queries."""
    result = FileParseResult(customer_id=config.customer.customer_id,
                             file_name=Path(source).name if isinstance(source,(str,Path)) else (file_name or ""))
    try:
        config = CustomerConfiguration.model_validate(config.model_dump(mode="json"))
        if not config.customer.enabled or not config.rule.enabled:
            raise FileProblem("CONFIG_ERROR", "客户配置或规则已禁用")
        if config.customer.sheet_name in {"字段映射", "字段映射说明"}:
            raise FileProblem("CONFIG_ERROR", "辅助字段映射说明工作表不能作为采购数据")
        result.config_version, result.rule_version = config.customer.config_version, config.rule.rule_version
        limits = limits or ParseLimits()
        name, content = _read_source(source, file_name, limits)
        result.file_name = name
        _parse_workbooks(config, content, result, limits, trust_formula_cache)
    except FileProblem as exc:
        result.file_error = exc.issue
    except ValidationError:
        result.file_error = ParseIssue(code="CONFIG_ERROR", message="配置字段或映射不合法")
    except (OSError, BadZipFile, InvalidFileException, ParseError, ZlibError, KeyError, ValueError, TypeError, IndexError, OverflowError, NotImplementedError) as exc:
        logger.warning("XLSX rejected: %s", type(exc).__name__)
        result.file_error = ParseIssue(code="INVALID_FILE", message="文件无法读取或不是有效的 XLSX 工作簿")
    if result.file_error is not None:
        result.rows.clear()  # No partial success for a corrupt or structurally invalid file.
    return result

def parse_purchase_file(session: Session, customer_id: str, source: Path | str | bytes, **options) -> FileParseResult:
    """Read module 02 configuration, then pure parse; does not query sales_order."""
    try:
        with session.no_autoflush:
            config = load_configuration(session, customer_id)
    except (ConfigurationError, ValidationError):
        return FileParseResult(customer_id=customer_id, file_name=options.get("file_name") or (Path(source).name if isinstance(source,(str,Path)) else ""),
                               file_error=ParseIssue(code="CONFIG_ERROR", message="客户配置缺失、禁用或不合法"))
    except SQLAlchemyError:
        logger.error("Cannot read customer configuration")
        return FileParseResult(customer_id=customer_id, file_name=options.get("file_name") or "",
                               file_error=ParseIssue(code="CONFIG_ERROR", message="数据库客户配置暂时无法读取"))
    return parse_excel(config, source, **options)
