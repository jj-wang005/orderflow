"""Read-only header metadata lookup. Old version receipts are never rewritten."""
from pathlib import Path
from io import BytesIO
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl import load_workbook
from sqlalchemy import select
from app.models import ConfigurationVersion

def configuration_columns(session, customer):
    version = session.scalar(select(ConfigurationVersion).where(
        ConfigurationVersion.customer_id == customer.customer_id,
        ConfigurationVersion.config_version == customer.config_version))
    stored = (version.snapshot or {}).get('source_columns') if version else None
    if stored:
        return stored, 'saved'
    # Only the unchanged built-in ABC definitions have a verifiable local source.
    # A custom configuration (including xiaomi) cannot borrow somebody else's header.
    from app.services.customer_configuration import read_configuration_file
    try:
        baseline = next((c.customer for c in read_configuration_file()
                         if c.customer.customer_id == customer.customer_id), None)
        if baseline and (customer.sheet_name, customer.header_row, customer.field_mapping) == (
                baseline.sheet_name, baseline.header_row, baseline.field_mapping):
            path = Path(__file__).resolve().parents[2] / 'sample_files' / (customer.customer_id+'客户采购订单模板.xlsx')
            if path.is_file():
                book = load_workbook(BytesIO(path.read_bytes()), read_only=True, data_only=False, keep_links=False)
                try:
                    sheet = book[customer.sheet_name]
                    sheet.reset_dimensions()
                    cells = list(next(sheet.iter_rows(min_row=customer.header_row,
                        max_row=customer.header_row, max_col=257)))
                    while cells and cells[-1].value is None:
                        cells.pop()
                    values = [cell.value for cell in cells]
                    if len(values) <= 256 and all(v is None or isinstance(v,str) for v in values) and all(
                            values.count(name)==1 for name in customer.field_mapping) and not any(cell.data_type=='f' for cell in cells):
                        return [{'index':i, 'name':v} for i,v in enumerate(values,1)], 'original_template'
                finally:
                    book.close()
    except (OSError, ValueError, TypeError, KeyError, StopIteration, BadZipFile, ParseError):
        # Recovery is optional metadata, never a reason to rewrite old receipts
        # or make a valid database configuration depend on an unavailable file.
        pass
    return [], 'upload_required'
