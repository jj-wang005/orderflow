"""92/93 focused tests: all writes rollback or use a disposable schema."""
import os
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
import fitz,pytest
from openpyxl import load_workbook
from sqlalchemy import text,select,func
from app.models import SalesOrder,OrderChange
from tests.shared_fixtures import empty_api as api
from tests.test_v2_batch_import import data,configure,seed,upload,rows,snapshot
from tests.concurrency_fixtures import client as concurrent_client, isolated_engine
from tests.test_v2_batch_concurrency import content
from tests.test_ai_drafts import create,edited,confirm
from tests.test_order_operations import edit_body

pytestmark=pytest.mark.skipif(os.getenv('RUN_DB_TESTS')!='1',reason='Explicit isolated MySQL only')

def made(c,conn):
 configure(c);seed(conn);blob=__import__('tests.test_v2_batch_import',fromlist=['excel']).excel()
 task=data(upload(c,blob),201);r=rows(c,task)[0];o=data(c.get('/api/v1/orders/'+str(r['created_order_id'])))
 return blob,task,r,o

def void(c,o,body=None):
 return c.post(f"/api/v1/orders/{o['id']}/void",json=body or {'request_key':str(uuid4()),'expected_revision':o['edit_revision'],'confirmed':True,'note':'测试：采购取消，保留来源'})

def test_void_terminal_audit_replay_and_duplicate_protection(api):
 c,conn=api;blob,task,row,o=made(c,conn);body={'request_key':str(uuid4()),'expected_revision':o['edit_revision'],'confirmed':True,'note':'采购取消'}
 claim=conn.execute(select(SalesOrder.purchase_source_digest).where(SalesOrder.id==o['id'])).scalar_one()
 saved=data(void(c,o,body))['order'];assert saved['status']=='VOIDED' and saved['edit_revision']==1
 assert data(void(c,o,body))['replayed']
 assert void(c,saved).status_code==409
 assert void(c,o,{**body,'note':'不同原因'}).status_code==409
 assert c.patch(f"/api/v1/orders/{o['id']}",json=edit_body(saved)).status_code==409
 assert c.patch(f"/api/v1/orders/{o['id']}/status",json={'request_key':str(uuid4()),'expected_revision':1,'status':'VALID'}).status_code==409
 logs=data(c.get(f"/api/v1/orders/{o['id']}/edit"))['history'];v=[x for x in logs if x['target'].startswith('void:')]
 assert len(v)==1 and v[0]['note']=='采购取消' and v[0]['before_data']['status']=='CREATED' and v[0]['after_data']['status']=='VOIDED'
 assert conn.scalar(select(SalesOrder.purchase_source_digest).where(SalesOrder.id==o['id']))==claim
 again=data(upload(c,blob,confirm=True),201);r=rows(c,again)[0]
 assert r['status']=='REVIEW_REQUIRED' and r['duplicate_order_id'] is None and not r['created_order_id']
 assert data(c.get('/api/v1/orders',params={'customer_id':'D'}))['total']==1
 assert data(c.get('/api/v1/orders',params={'customer_id':'D','status':'VOIDED'}))['items'][0]['id']==o['id']

@pytest.mark.parametrize('status',['VALID','PROCESSING','SHIPPED','COMPLETED','INVALID','VOIDED'])
def test_non_created_void_refused(api,status):
 c,conn=api;_,_,_,o=made(c,conn);conn.execute(SalesOrder.__table__.update().where(SalesOrder.id==o['id']).values(status=status));before=snapshot(conn)
 assert void(c,o).status_code==409 and snapshot(conn)==before

@pytest.mark.parametrize('delta',[{'note':'  '},{'confirmed':False},{'expected_revision':99}])
def test_void_requires_reason_confirmation_and_revision(api,delta):
 c,conn=api;_,_,_,o=made(c,conn)
 body={'request_key':str(uuid4()),'expected_revision':0,'confirmed':True,'note':'取消',**delta};before=snapshot(conn)
 assert void(c,o,body).status_code in (409,422) and snapshot(conn)==before

def test_seed_and_unknown_void_refused(api):
 c,conn=api;configure(c);oid=seed(conn)
 assert void(c,{'id':oid,'edit_revision':0}).status_code==409
 assert void(c,{'id':999999999,'edit_revision':0}).status_code==404

def test_void_not_eligible_history_and_manual_provenance(api):
 from app.services.matching_engine import match_purchase_row
 from app.schemas.purchase_row import NormalizedPurchaseRow
 c,conn=api;_,task,row,o=made(c,conn);data(void(c,o))
 # Original valid SEED removed only inside this test transaction, leaving void alone.
 conn.execute(SalesOrder.__table__.update().where(SalesOrder.customer_id=='D',SalesOrder.source_type=='SEED').values(status='INVALID'))
 from sqlalchemy.orm import Session
 with Session(bind=conn,join_transaction_mode='create_savepoint') as s:
  import json
  match=match_purchase_row(s,NormalizedPurchaseRow.model_validate_json(json.dumps(row['normalized_data'])))
  assert match.status=='NOT_MATCHED'

def test_excel_all_filtered_pages_text_and_amount(api,monkeypatch):
 c,conn=api;configure(c)
 from decimal import Decimal
 from datetime import datetime
 from sqlalchemy.orm import Session
 with Session(bind=conn,join_transaction_mode='create_savepoint') as s,s.begin():
  for i in range(23):s.add(SalesOrder(order_no='EXPORT-'+str(i),customer_id='D',customer_order_no='=HYPERLINK("bad")',customer_material_code='+cmd',customer_model='@SUM(1)',customer_batch='-1+1',product_code='P',product_name='中文产品',quantity=Decimal(3),unit_price=Decimal('1.2345'),amount=Decimal('3.70'),status='VALID',source_type='SEED',created_at=datetime(2026,1,2)))
 filters={'customer_id':'D','status':'VALID','date_from':'2026-01-02','date_to':'2026-01-02'}
 page=data(c.get('/api/v1/orders',params={**filters,'page_size':1}));assert page['total']==23
 before=snapshot(conn);response=c.get('/api/v1/orders/export.xlsx',params=filters);assert response.status_code==200
 book=load_workbook(BytesIO(response.content));sheet=book.active
 assert sheet.max_row==24 and sheet['D2'].value.startswith('=') and sheet['D2'].data_type=='s'
 assert all(sheet.cell(2,col).data_type=='s' for col in (4,5,6,7))
 assert sheet['L2'].value==3.7 and sheet['L2'].number_format=='#,##0.00'
 assert sheet['A2'].value==page['items'][0]['order_no'] and snapshot(conn)==before
 assert c.get('/api/v1/orders/export.xlsx',params={**filters,'date_to':'9999-12-31'}).status_code==200
 assert c.get('/api/v1/orders/export.xlsx',params={**filters,'date_to':'2026-01-01'}).status_code==422
 import app.services.order_exports as exports
 monkeypatch.setattr(exports,'MAX_EXPORT_ROWS',10)
 assert c.get('/api/v1/orders/export.xlsx',params=filters).status_code==422

def test_void_excel_filter_and_chinese_pdf_report(api, tmp_path):
 c,conn=api;_,_,_,o=made(c,conn);out=tmp_path
 conn.execute(SalesOrder.__table__.update().where(SalesOrder.id==o['id']).values(product_name='长产品名称'*40))
 normal=c.get(f"/api/v1/orders/{o['id']}/report.pdf");assert normal.status_code==200
 doc=fitz.open(stream=normal.content,filetype='pdf');assert '内部销售订单详情' in ''.join(p.get_text() for p in doc);doc.close()
 (out/'sample_normal.pdf').write_bytes(normal.content)
 long_reason='测试：采购取消。'+'已核实来源，保留原采购及作废审计。'*28
 data(void(c,o,{'request_key':str(uuid4()),'expected_revision':0,'confirmed':True,'note':long_reason}));before=snapshot(conn)
 response=c.get(f"/api/v1/orders/{o['id']}/report.pdf");assert response.status_code==200
 doc=fitz.open(stream=response.content,filetype='pdf');text_all=''.join(p.get_text() for p in doc)
 assert '已作废' in text_all and '采购取消' in text_all and '后端保存总金额' in text_all and '采购行ID' in text_all
 assert len(doc)>=2 and text_all.replace('\n','').count('已核实来源，保留原采购及作废审计。')>=20
 assert all(word[2]<=554 and word[3]<=830 for p in doc for word in p.get_text('words'))
 assert 'MIMO_API_KEY' not in text_all and 'DB_PASSWORD' not in text_all
 (out/'sample_voided.pdf').write_bytes(response.content)
 for n,p in enumerate(doc):p.get_pixmap(matrix=fitz.Matrix(1.25,1.25)).save(out/f'pdf_voided_{n+1}.png')
 doc.close();assert snapshot(conn)==before
 filtered=load_workbook(BytesIO(c.get('/api/v1/orders/export.xlsx',params={'customer_id':'D','status':'VOIDED'}).content)).active
 assert filtered.max_row==2 and filtered['M2'].value=='已作废'
 default=load_workbook(BytesIO(c.get('/api/v1/orders/export.xlsx',params={'customer_id':'D'}).content)).active
 assert all(default.cell(i,13).value!='已作废' for i in range(2,default.max_row+1))

def test_ai_incomplete_invalid_draft_saves_but_cannot_confirm(api):
 c,conn=api;configure(c);d=create(c)
 required=d['required_fields'];assert required==['customer_order_no','customer_material_code','quantity']
 result=d['edited_result'];result['rows'][0]['quantity']='不是数量';result['rows'][0]['customer_order_no']=None
 edited_d=data(c.put(f"/api/v1/ai-extractions/{d['id']}",json={'request_key':str(uuid4()),'expected_revision':d['revision'],'note':'保存未完成草稿','result':result}))
 assert edited_d['edited_result']['rows'][0]['quantity']=='不是数量'
 before=snapshot(conn);assert confirm(c,edited_d).status_code==422 and snapshot(conn)==before

def test_ai_removed_rows_preserved_and_not_imported(api):
 c,conn=api;configure(c);d=create(c);result=d['edited_result'];removed=result['rows'].pop();result['removed_rows']=[removed]
 result['rows'][0].update(customer_order_no='PO-NEW',customer_material_code='MAT',quantity='12',uncertain_fields=[])
 d=data(c.put(f"/api/v1/ai-extractions/{d['id']}",json={'request_key':str(uuid4()),'expected_revision':d['revision'],'note':'移除重复识别','result':result}))
 confirmed=data(confirm(c,d));task=data(c.get('/api/v1/imports/'+str(confirmed['import_task_id'])))
 assert task['total_rows']==1 and len(confirmed['edited_result']['removed_rows'])==1 and len(confirmed['original_result']['rows'])==2

def test_match_explanation_is_snapshot_readonly(api):
 c,conn=api;configure(c);task=data(upload(c),201);row=rows(c,task)[0];before=snapshot(conn)
 result=data(c.get(f"/api/v1/imports/{task['task_id']}/rows/{row['id']}/match-explanation"))
 assert result['current_result']['status']=='NOT_MATCHED' and result['config_version']==task['config_version']
 assert result['match_values']['customer_order_no']=='PO-01' and snapshot(conn)==before

def test_task_presentation_separates_file_failure_from_unmatched(api):
 c,conn=api;configure(c)
 parsed=data(upload(c),201);failed=data(upload(c,b'broken'),201)
 assert parsed['status_description']=='已解析 · 待人工补全'
 assert rows(c,parsed)[0]['status_description']=='已解析 · 待人工补全'
 assert failed['status_description']=='文件解析失败' and failed['total_rows']==0

def test_a_rule_check_includes_material_safety_field(abc_api):
 c,conn=abc_api
 content=Path('sample_files/A客户采购订单模板.xlsx').read_bytes()
 task=data(upload(c,content,cid='A',confirm=True),201);row=rows(c,task)[0]
 result=data(c.get(f"/api/v1/imports/{task['task_id']}/rows/{row['id']}/match-explanation"))
 assert 'customer_order_no' in result['match_fields'] and 'customer_material_code' in result['match_fields']
 assert result['match_values']['customer_material_code']==row['normalized_data']['customer_material_code']

@pytest.mark.parametrize('other',['status','different-void','same-void'])
def test_concurrent_void_and_status_only_one_wins(concurrent_client,other):
 c,engine=concurrent_client
 task=data(upload(c,content()),201);row=rows(c,task)[0];o=data(c.get('/api/v1/orders/'+str(row['created_order_id'])))
 barrier=Barrier(2);body={'request_key':str(uuid4()),'expected_revision':0,'confirmed':True,'note':'并发取消'}
 def run(which):
  barrier.wait(timeout=10)
  if which=='void':return void(c,o,body)
  if other=='same-void':return void(c,o,body)
  if other=='different-void':return void(c,o)
  return c.patch(f"/api/v1/orders/{o['id']}/status",json={'request_key':str(uuid4()),'expected_revision':0,'status':'VALID'})
 with ThreadPoolExecutor(max_workers=2) as pool:responses=list(pool.map(run,['void','status']))
 assert sorted(r.status_code for r in responses)==([200,200] if other=='same-void' else [200,409])
 with engine.connect() as conn:
  assert conn.scalar(select(func.count()).select_from(OrderChange).where(OrderChange.order_id==o['id']))==1
  assert conn.scalar(select(SalesOrder.edit_revision).where(SalesOrder.id==o['id']))==1

from tests.shared_fixtures import abc_api
